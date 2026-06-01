"""Build sponsor_scores.json from OFLC disclosure CSV/XLSX files.

Download H-1B/H-1B1/E-3 disclosure data from:
https://www.dol.gov/agencies/eta/foreign-labor/performance

Usage:
    python scripts/build_sponsor_scores.py path/to/H-1B_Disclosure.csv
    python scripts/build_sponsor_scores.py path/to/file1.csv path/to/file2.xlsx
"""

import argparse
import csv
import json
import os
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from company_registry import normalize_company_name


DEFAULT_OUTPUT = os.path.join(
    os.path.dirname(os.path.dirname(__file__)), "data", "sponsor_scores.json"
)

EMPLOYER_FIELDS = ("EMPLOYER_NAME", "Employer Name", "employer_name")
STATUS_FIELDS = ("CASE_STATUS", "Case Status", "case_status")
VISA_FIELDS = ("VISA_CLASS", "Visa Class", "visa_class")

SUPPORTED_VISA_CLASSES = {"H-1B", "H-1B1 Chile", "H-1B1 Singapore", "E-3 Australian"}
CERTIFIED_STATUSES = {"Certified", "Certified - Withdrawn"}


def _first_value(row, names):
    for name in names:
        value = row.get(name)
        if value:
            return str(value).strip()
    return ""


def _rows_from_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        yield from csv.DictReader(f)


def _rows_from_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "Reading .xlsx files requires openpyxl. Install it or export the OFLC file as CSV."
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    for values in rows:
        yield {
            headers[idx]: value
            for idx, value in enumerate(values)
            if idx < len(headers)
        }


def _iter_rows(path):
    lower = path.lower()
    if lower.endswith(".csv"):
        yield from _rows_from_csv(path)
    elif lower.endswith(".xlsx"):
        yield from _rows_from_xlsx(path)
    else:
        raise SystemExit(f"Unsupported file type: {path}")


def build_scores(paths):
    totals = Counter()
    statuses = defaultdict(Counter)
    display_names = {}

    for path in paths:
        for row in _iter_rows(path):
            employer = _first_value(row, EMPLOYER_FIELDS)
            if not employer:
                continue
            visa_class = _first_value(row, VISA_FIELDS)
            if visa_class and visa_class not in SUPPORTED_VISA_CLASSES:
                continue
            status = _first_value(row, STATUS_FIELDS) or "Unknown"
            key = normalize_company_name(employer)
            display_names.setdefault(key, employer)
            statuses[key][status] += 1
            if status in CERTIFIED_STATUSES:
                totals[key] += 1

    return {
        key: {
            "display_name": display_names[key],
            "recent_lca_count": totals[key],
            "statuses": dict(statuses[key]),
        }
        for key in sorted(display_names)
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+", help="OFLC disclosure CSV/XLSX files")
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    scores = build_scores(args.files)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(scores, f, indent=2, sort_keys=True)
        f.write("\n")
    print(f"Wrote {len(scores)} sponsor score(s) to {args.output}")


if __name__ == "__main__":
    main()
